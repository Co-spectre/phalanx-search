"""
Phalanx Search - Document Parsers
Extracts text content from 30+ document formats.
100% Local processing — No data leaves your system.
"""

import os
import re
import json
import email
import sqlite3
from pathlib import Path
from typing import Dict, List, Optional
from dataclasses import dataclass
from datetime import datetime
import hashlib


@dataclass
class ParsedDocument:
    """Represents a parsed document with its content and metadata"""
    filename: str
    filepath: str
    content: str
    file_type: str
    file_size: int
    page_count: int
    created_at: datetime
    doc_hash: str
    metadata: Dict


class BaseParser:
    """Base class for document parsers"""

    def parse(self, filepath: str) -> ParsedDocument:
        raise NotImplementedError

    def _get_file_hash(self, filepath: str) -> str:
        """Generate MD5 hash of file for deduplication"""
        hasher = hashlib.md5()
        with open(filepath, 'rb') as f:
            buf = f.read(65536)
            while len(buf) > 0:
                hasher.update(buf)
                buf = f.read(65536)
        return hasher.hexdigest()

    def _get_file_info(self, filepath: str) -> Dict:
        """Get basic file information"""
        path = Path(filepath)
        stat = path.stat()
        return {
            "filename": path.name,
            "filepath": str(path.absolute()),
            "file_size": stat.st_size,
            "file_type": path.suffix.lower().replace(".", ""),
            "created_at": datetime.fromtimestamp(stat.st_ctime),
            "doc_hash": self._get_file_hash(filepath)
        }

    def _read_text_file(self, filepath: str) -> str:
        """Read text file with encoding fallback."""
        encodings = ['utf-8', 'utf-8-sig', 'utf-16', 'latin-1', 'cp1252', 'ascii']
        for encoding in encodings:
            try:
                with open(filepath, 'r', encoding=encoding) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        # Last resort: read as bytes and decode with replacement
        with open(filepath, 'rb') as f:
            return f.read().decode('utf-8', errors='replace')


# ============== Core Document Parsers ==============

class PDFParser(BaseParser):
    """Parser for PDF documents using PyMuPDF"""

    def parse(self, filepath: str) -> ParsedDocument:
        import fitz  # PyMuPDF
        file_info = self._get_file_info(filepath)

        doc = fitz.open(filepath)
        text_content = []

        for page_num, page in enumerate(doc):
            text = page.get_text()
            if text.strip():
                text_content.append(f"[Page {page_num + 1}]\n{text}")

        page_count = len(doc)
        doc.close()

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content="\n\n".join(text_content),
            file_type="pdf",
            file_size=file_info["file_size"],
            page_count=page_count,
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "pdf", "pages": page_count}
        )


class WordParser(BaseParser):
    """Parser for Word documents (.docx)"""

    def parse(self, filepath: str) -> ParsedDocument:
        from docx import Document
        file_info = self._get_file_info(filepath)

        doc = Document(filepath)
        text_content = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_content.append(para.text)

        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                table_text.append(" | ".join(row_text))
            if table_text:
                text_content.append("\n[Table]\n" + "\n".join(table_text))

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content="\n\n".join(text_content),
            file_type="docx",
            file_size=file_info["file_size"],
            page_count=len(doc.paragraphs) // 30 + 1,
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "word", "paragraphs": len(doc.paragraphs)}
        )


class ExcelParser(BaseParser):
    """Parser for Excel spreadsheets (.xlsx)"""

    def parse(self, filepath: str) -> ParsedDocument:
        import pandas as pd
        file_info = self._get_file_info(filepath)

        text_content = []
        sheet_count = 0

        try:
            xlsx = pd.ExcelFile(filepath)
            sheet_names = xlsx.sheet_names
            sheet_count = len(sheet_names)

            for sheet_name in sheet_names:
                df = pd.read_excel(xlsx, sheet_name=sheet_name)
                sheet_text = f"[Sheet: {sheet_name}]\n"
                sheet_text += df.to_string(index=False, max_rows=500)
                text_content.append(sheet_text)

            xlsx.close()
        except Exception as e:
            text_content.append(f"Error reading Excel: {e}")

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content="\n\n".join(text_content),
            file_type="xlsx",
            file_size=file_info["file_size"],
            page_count=sheet_count,
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "excel", "sheets": sheet_count}
        )


class PowerPointParser(BaseParser):
    """Parser for PowerPoint presentations (.pptx)"""

    def parse(self, filepath: str) -> ParsedDocument:
        from pptx import Presentation
        file_info = self._get_file_info(filepath)

        prs = Presentation(filepath)
        text_content = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = [f"[Slide {slide_num}]"]
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
            if len(slide_text) > 1:
                text_content.append("\n".join(slide_text))

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content="\n\n".join(text_content),
            file_type="pptx",
            file_size=file_info["file_size"],
            page_count=len(prs.slides),
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "powerpoint", "slides": len(prs.slides)}
        )


class TextParser(BaseParser):
    """Parser for plain text files (.txt, .md, .csv, .log, .ini, .cfg, .conf, .yaml, .yml, .toml)"""

    def parse(self, filepath: str) -> ParsedDocument:
        file_info = self._get_file_info(filepath)
        content = self._read_text_file(filepath)
        page_count = max(1, len(content) // 3000)

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content=content,
            file_type=file_info["file_type"],
            file_size=file_info["file_size"],
            page_count=page_count,
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "text", "chars": len(content)}
        )


# ============== Extended Parsers ==============

class HTMLParser(BaseParser):
    """Parser for HTML files — extracts readable text content."""

    def parse(self, filepath: str) -> ParsedDocument:
        from bs4 import BeautifulSoup
        file_info = self._get_file_info(filepath)

        raw = self._read_text_file(filepath)
        soup = BeautifulSoup(raw, "lxml" if "lxml" in _available_modules() else "html.parser")

        # Remove script and style elements
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()

        # Extract text with structure
        text_parts = []

        # Title
        title = soup.find("title")
        if title and title.string:
            text_parts.append(f"Title: {title.string.strip()}")

        # Headings and paragraphs
        for elem in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th", "pre", "code"]):
            text = elem.get_text(separator=" ", strip=True)
            if text and len(text) > 2:
                tag = elem.name
                if tag.startswith("h"):
                    text_parts.append(f"\n[{tag.upper()}] {text}")
                else:
                    text_parts.append(text)

        content = "\n".join(text_parts) if text_parts else soup.get_text(separator="\n", strip=True)

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content=content,
            file_type="html",
            file_size=file_info["file_size"],
            page_count=1,
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "html"}
        )


class JSONParser(BaseParser):
    """Parser for JSON and JSONL files — flattens nested structures into searchable text."""

    def parse(self, filepath: str) -> ParsedDocument:
        file_info = self._get_file_info(filepath)
        raw = self._read_text_file(filepath)
        ext = Path(filepath).suffix.lower()

        text_parts = []

        try:
            if ext == ".jsonl":
                # JSON Lines: one JSON object per line
                for i, line in enumerate(raw.strip().split("\n")):
                    if line.strip():
                        try:
                            obj = json.loads(line)
                            text_parts.append(f"[Record {i + 1}]\n{self._flatten_json(obj)}")
                        except json.JSONDecodeError:
                            continue
            else:
                data = json.loads(raw)
                if isinstance(data, list):
                    for i, item in enumerate(data[:500]):  # Limit for huge arrays
                        text_parts.append(f"[Item {i + 1}]\n{self._flatten_json(item)}")
                else:
                    text_parts.append(self._flatten_json(data))
        except json.JSONDecodeError:
            text_parts.append(raw[:10000])  # Fallback to raw text

        content = "\n\n".join(text_parts)

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content=content,
            file_type=file_info["file_type"],
            file_size=file_info["file_size"],
            page_count=1,
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "json"}
        )

    def _flatten_json(self, obj, prefix: str = "") -> str:
        """Recursively flatten a JSON object into readable key: value lines."""
        lines = []
        if isinstance(obj, dict):
            for key, value in obj.items():
                full_key = f"{prefix}.{key}" if prefix else key
                if isinstance(value, (dict, list)):
                    lines.append(self._flatten_json(value, full_key))
                else:
                    lines.append(f"{full_key}: {value}")
        elif isinstance(obj, list):
            for i, item in enumerate(obj[:100]):
                lines.append(self._flatten_json(item, f"{prefix}[{i}]"))
        else:
            lines.append(f"{prefix}: {obj}" if prefix else str(obj))
        return "\n".join(lines)


class XMLParser(BaseParser):
    """Parser for XML files — extracts text content preserving structure."""

    def parse(self, filepath: str) -> ParsedDocument:
        import xml.etree.ElementTree as ET
        file_info = self._get_file_info(filepath)

        text_parts = []
        try:
            tree = ET.parse(filepath)
            root = tree.getroot()
            self._extract_xml_text(root, text_parts)
        except ET.ParseError:
            # Fallback to raw text
            content = self._read_text_file(filepath)
            # Strip XML tags manually
            text_parts.append(re.sub(r'<[^>]+>', ' ', content))

        content = "\n".join(text_parts)

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content=content,
            file_type="xml",
            file_size=file_info["file_size"],
            page_count=1,
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "xml"}
        )

    def _extract_xml_text(self, element, text_parts: list, depth: int = 0):
        """Recursively extract text from XML elements."""
        # Strip namespace from tag
        tag = re.sub(r'\{[^}]+\}', '', element.tag)

        if element.text and element.text.strip():
            text_parts.append(f"{'  ' * depth}{tag}: {element.text.strip()}")

        for child in element:
            self._extract_xml_text(child, text_parts, depth + 1)

        if element.tail and element.tail.strip():
            text_parts.append(element.tail.strip())


class SQLiteParser(BaseParser):
    """Parser for SQLite database files — extracts schema and data as searchable text."""

    def parse(self, filepath: str) -> ParsedDocument:
        file_info = self._get_file_info(filepath)
        text_parts = []

        try:
            conn = sqlite3.connect(f"file:{filepath}?mode=ro", uri=True)
            conn.text_factory = lambda b: b.decode('utf-8', errors='replace')

            # Get all tables
            tables = conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
            ).fetchall()

            for (table_name,) in tables:
                text_parts.append(f"\n[Table: {table_name}]")

                # Get schema
                try:
                    schema = conn.execute(f"PRAGMA table_info('{table_name}')").fetchall()
                    columns = [row[1] for row in schema]
                    text_parts.append(f"Columns: {', '.join(columns)}")

                    # Get row count
                    count = conn.execute(f"SELECT COUNT(*) FROM '{table_name}'").fetchone()[0]
                    text_parts.append(f"Rows: {count}")

                    # Extract sample data (up to 200 rows)
                    rows = conn.execute(f"SELECT * FROM '{table_name}' LIMIT 200").fetchall()
                    for row in rows:
                        row_text = " | ".join(
                            f"{col}: {val}" for col, val in zip(columns, row)
                            if val is not None and str(val).strip()
                        )
                        if row_text:
                            text_parts.append(row_text)
                except Exception as e:
                    text_parts.append(f"Error reading table {table_name}: {e}")

            conn.close()
        except Exception as e:
            text_parts.append(f"Error reading SQLite database: {e}")

        content = "\n".join(text_parts)

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content=content,
            file_type="sqlite",
            file_size=file_info["file_size"],
            page_count=len([t for t in text_parts if t.startswith("\n[Table:")]),
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "sqlite"}
        )


class EPUBParser(BaseParser):
    """Parser for EPUB e-books."""

    def parse(self, filepath: str) -> ParsedDocument:
        file_info = self._get_file_info(filepath)
        text_parts = []

        try:
            import ebooklib
            from ebooklib import epub
            from bs4 import BeautifulSoup

            book = epub.read_epub(filepath, options={"ignore_ncx": True})

            # Extract title
            title = book.get_metadata('DC', 'title')
            if title:
                text_parts.append(f"Title: {title[0][0]}")

            # Extract content from HTML items
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                soup = BeautifulSoup(item.get_content(), "html.parser")
                text = soup.get_text(separator="\n", strip=True)
                if text and len(text) > 10:
                    text_parts.append(text)

        except ImportError:
            text_parts.append("EPUB parsing requires ebooklib: pip install ebooklib")
        except Exception as e:
            text_parts.append(f"Error reading EPUB: {e}")

        content = "\n\n".join(text_parts)

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content=content,
            file_type="epub",
            file_size=file_info["file_size"],
            page_count=max(1, len(text_parts)),
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "epub"}
        )


class EmailParser(BaseParser):
    """Parser for email files (.eml, .msg)."""

    def parse(self, filepath: str) -> ParsedDocument:
        file_info = self._get_file_info(filepath)
        ext = Path(filepath).suffix.lower()

        if ext == ".msg":
            return self._parse_msg(filepath, file_info)
        return self._parse_eml(filepath, file_info)

    def _parse_eml(self, filepath: str, file_info: Dict) -> ParsedDocument:
        """Parse .eml files using stdlib email module."""
        text_parts = []
        try:
            with open(filepath, 'rb') as f:
                msg = email.message_from_binary_file(f)

            # Headers
            for header in ["From", "To", "Subject", "Date"]:
                value = msg.get(header)
                if value:
                    text_parts.append(f"{header}: {value}")

            # Body
            if msg.is_multipart():
                for part in msg.walk():
                    if part.get_content_type() == "text/plain":
                        payload = part.get_payload(decode=True)
                        if payload:
                            text_parts.append(payload.decode('utf-8', errors='replace'))
            else:
                payload = msg.get_payload(decode=True)
                if payload:
                    text_parts.append(payload.decode('utf-8', errors='replace'))

        except Exception as e:
            text_parts.append(f"Error reading email: {e}")

        content = "\n\n".join(text_parts)
        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content=content,
            file_type="eml",
            file_size=file_info["file_size"],
            page_count=1,
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "email"}
        )

    def _parse_msg(self, filepath: str, file_info: Dict) -> ParsedDocument:
        """Parse .msg (Outlook) files."""
        text_parts = []
        try:
            import extract_msg
            msg = extract_msg.Message(filepath)

            if msg.subject:
                text_parts.append(f"Subject: {msg.subject}")
            if msg.sender:
                text_parts.append(f"From: {msg.sender}")
            if msg.to:
                text_parts.append(f"To: {msg.to}")
            if msg.date:
                text_parts.append(f"Date: {msg.date}")
            if msg.body:
                text_parts.append(msg.body)

            msg.close()
        except ImportError:
            text_parts.append("MSG parsing requires extract-msg: pip install extract-msg")
        except Exception as e:
            text_parts.append(f"Error reading MSG: {e}")

        content = "\n\n".join(text_parts)
        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content=content,
            file_type="msg",
            file_size=file_info["file_size"],
            page_count=1,
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "email_msg"}
        )


class CodeParser(BaseParser):
    """Parser for source code files — preserves structure and adds context."""

    # Language identifiers for syntax context
    LANG_MAP = {
        ".py": "Python", ".js": "JavaScript", ".ts": "TypeScript",
        ".java": "Java", ".cpp": "C++", ".c": "C", ".h": "C/C++ Header",
        ".cs": "C#", ".go": "Go", ".rs": "Rust", ".rb": "Ruby",
        ".php": "PHP", ".swift": "Swift", ".kt": "Kotlin", ".r": "R",
        ".sql": "SQL", ".sh": "Shell", ".bat": "Batch", ".ps1": "PowerShell",
    }

    def parse(self, filepath: str) -> ParsedDocument:
        file_info = self._get_file_info(filepath)
        content = self._read_text_file(filepath)
        ext = Path(filepath).suffix.lower()
        lang = self.LANG_MAP.get(ext, "Code")

        # Add language context header
        header = f"[{lang} Source File: {file_info['filename']}]\n"

        # Extract structural elements (functions, classes, etc.)
        structures = self._extract_structures(content, ext)
        if structures:
            header += f"Structures: {', '.join(structures)}\n\n"

        full_content = header + content

        return ParsedDocument(
            filename=file_info["filename"],
            filepath=file_info["filepath"],
            content=full_content,
            file_type=file_info["file_type"],
            file_size=file_info["file_size"],
            page_count=max(1, content.count('\n') // 50),
            created_at=file_info["created_at"],
            doc_hash=file_info["doc_hash"],
            metadata={"source": "code", "language": lang}
        )

    def _extract_structures(self, content: str, ext: str) -> List[str]:
        """Extract function/class names from code."""
        structures = []
        patterns = {
            ".py": [r'(?:def|class)\s+(\w+)', r'async\s+def\s+(\w+)'],
            ".js": [r'(?:function|class)\s+(\w+)', r'(?:const|let|var)\s+(\w+)\s*=\s*(?:async\s+)?(?:\([^)]*\)\s*=>|function)'],
            ".ts": [r'(?:function|class|interface|type|enum)\s+(\w+)'],
            ".java": [r'(?:class|interface|enum)\s+(\w+)', r'(?:public|private|protected)\s+\w+\s+(\w+)\s*\('],
            ".cpp": [r'(?:class|struct)\s+(\w+)', r'\w+\s+(\w+)\s*\([^)]*\)\s*\{'],
            ".c": [r'\w+\s+(\w+)\s*\([^)]*\)\s*\{'],
            ".go": [r'func\s+(?:\(\w+\s+\*?\w+\)\s+)?(\w+)'],
            ".rs": [r'(?:fn|struct|enum|trait|impl)\s+(\w+)'],
        }

        for pattern in patterns.get(ext, []):
            matches = re.findall(pattern, content)
            structures.extend(matches[:20])

        return structures[:30]


def _available_modules() -> set:
    """Check which optional modules are available."""
    available = set()
    for mod in ["lxml", "bs4", "ebooklib", "extract_msg"]:
        try:
            __import__(mod)
            available.add(mod)
        except ImportError:
            pass
    return available


# ============== Factory ==============

class DocumentParserFactory:
    """Factory to get the appropriate parser for a file type"""

    PARSERS = {
        # Core documents
        ".pdf": PDFParser,
        ".docx": WordParser,
        ".doc": WordParser,
        ".xlsx": ExcelParser,
        ".xls": ExcelParser,
        ".pptx": PowerPointParser,
        # Text files
        ".txt": TextParser,
        ".md": TextParser,
        ".csv": TextParser,
        ".rtf": TextParser,
        ".log": TextParser,
        ".ini": TextParser,
        ".cfg": TextParser,
        ".conf": TextParser,
        ".yaml": TextParser,
        ".yml": TextParser,
        ".toml": TextParser,
        # Web & Structured
        ".html": HTMLParser,
        ".htm": HTMLParser,
        ".xml": XMLParser,
        ".json": JSONParser,
        ".jsonl": JSONParser,
        # Databases
        ".sqlite": SQLiteParser,
        ".db": SQLiteParser,
        # E-books
        ".epub": EPUBParser,
        # Email
        ".eml": EmailParser,
        ".msg": EmailParser,
        # Code
        ".py": CodeParser,
        ".js": CodeParser,
        ".ts": CodeParser,
        ".java": CodeParser,
        ".cpp": CodeParser,
        ".c": CodeParser,
        ".h": CodeParser,
        ".cs": CodeParser,
        ".go": CodeParser,
        ".rs": CodeParser,
        ".rb": CodeParser,
        ".php": CodeParser,
        ".swift": CodeParser,
        ".kt": CodeParser,
        ".r": CodeParser,
        ".sql": CodeParser,
        ".sh": CodeParser,
        ".bat": CodeParser,
        ".ps1": CodeParser,
    }

    @classmethod
    def get_parser(cls, filepath: str) -> Optional[BaseParser]:
        ext = Path(filepath).suffix.lower()
        parser_class = cls.PARSERS.get(ext)
        if parser_class:
            return parser_class()
        return None

    @classmethod
    def parse_document(cls, filepath: str) -> Optional[ParsedDocument]:
        parser = cls.get_parser(filepath)
        if parser:
            return parser.parse(filepath)
        return None

    @classmethod
    def is_supported(cls, filepath: str) -> bool:
        ext = Path(filepath).suffix.lower()
        return ext in cls.PARSERS
